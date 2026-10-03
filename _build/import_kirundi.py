"""Load the Kirundi column of the translation template into app/i18n.js.

Usage:  python _build/import_kirundi.py "<filled template>.xlsx" [app/i18n.js]

- Reads sheet "Strings": column F (Kirundi) and column I (Key).
- Rebuilds the `rn` block of i18n.js. Text items left empty fall back to French in
  the app. Inside a group (result names, advice lists) a missing line is filled
  with the French line so a screen is never half broken; the report lists them.
- Reports advice rows not marked "Checked by vet = Yes".
"""
import json
import re
import sys
from openpyxl import load_workbook

xlsx = sys.argv[1]
i18n_path = sys.argv[2] if len(sys.argv) > 2 else "app/i18n.js"

ws = load_workbook(xlsx, data_only=True)["Strings"]
head = [str(c.value or "").strip() for c in ws[1]]
col = {name: head.index(name) for name in head}
K, RN, FR, VET = col["Key (do not change)"], col["Kirundi (to write)"], col["French"], col["Checked by vet"]

rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r[K]:
        continue
    rows.append({"key": str(r[K]).strip(), "rn": (str(r[RN]).strip() if r[RN] else ""),
                 "fr": (str(r[FR]).strip() if r[FR] else ""), "vet": str(r[VET] or "").strip()})

# rebuild nested structure: "advice.cocci.2" -> advice -> cocci -> [.., line 2]
def put(tree, parts, value):
    head, rest = parts[0], parts[1:]
    if not rest:
        tree[head] = value
        return
    tree.setdefault(head, {})
    put(tree[head], rest, value)

rn, fallback, done, unchecked = {}, [], 0, []
groups = {r["key"].split(".")[0] for r in rows if "." in r["key"]}
group_has_rn = {g: any(r["rn"] for r in rows if r["key"].startswith(g + ".")) for g in groups}
for r in rows:
    parts = r["key"].split(".")
    top = parts[0]
    if len(parts) == 1:
        if r["rn"]:
            rn[top] = r["rn"]; done += 1
        continue
    if not group_has_rn[top]:
        continue                              # whole group untranslated: app falls back to French
    if r["rn"]:
        done += 1
    else:
        fallback.append(r["key"])
    put(rn, parts, r["rn"] or r["fr"])
    if top == "advice" and r["rn"] and r["vet"].lower() != "yes":
        unchecked.append(r["key"])

# numbered leaves under advice.<class> become lists, in order
def listify(node):
    if isinstance(node, dict):
        if node and all(k.isdigit() for k in node):
            return [listify(node[k]) for k in sorted(node, key=int)]
        return {k: listify(v) for k, v in node.items()}
    return node
rn = listify(rn)

block = "rn: /*RN-START*/" + json.dumps(rn, ensure_ascii=False, indent=2).replace("\n", "\n  ") + "/*RN-END*/,"
src = open(i18n_path, encoding="utf-8").read()
if "/*RN-START*/" in src:
    src = re.sub(r"rn: /\*RN-START\*/.*?/\*RN-END\*/,", lambda m: block, src, flags=re.S)
elif re.search(r"\n  rn: \{\},", src):
    src = re.sub(r"\n  rn: \{\},", lambda m: "\n  " + block, src)
else:
    sys.exit("Could not find the rn block in " + i18n_path)
open(i18n_path, "w", encoding="utf-8", newline="\n").write(src)

total = len(rows)
print(f"Kirundi loaded: {done} of {total} strings.")
if fallback:
    print(f"Filled from French inside partly translated groups ({len(fallback)}): " + ", ".join(fallback))
missing_top = [r["key"] for r in rows if "." not in r["key"] and not r["rn"]]
if missing_top:
    print(f"Still in French ({len(missing_top)}): " + ", ".join(missing_top))
if unchecked:
    print(f"ADVICE NOT YET CHECKED BY THE VET ({len(unchecked)}): " + ", ".join(unchecked))
